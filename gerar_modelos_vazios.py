# -*- coding: utf-8 -*-
"""Gera as versões EM BRANCO (para download) das duas planilhas do Saiaê.

Parte dos arquivos já gerados (rode antes gerar_planilha.py e gerar_planilha_enxuta.py),
apaga todos os números e datas dos campos editáveis e ajusta os textos que falavam dos exemplos.
Mantém: estrutura, fórmulas, visual, nomes dos custos fixos e o cenário "Realista" selecionado.
O mês de início vira "próximo mês" (fórmula) até a pessoa digitar a data dela.
Rode:  python3 gerar_modelos_vazios.py
"""
from openpyxl import load_workbook

from estilo_app import INPUT

ARQUIVOS = {
    "Saiae_Financeiro.xlsx": "Saiae_Financeiro_Modelo.xlsx",
    "Saiae_Enxuta.xlsx": "Saiae_Enxuta_Modelo.xlsx",
}
TEXTOS = {
    "Só os PREÇOS dos planos vieram do escopo. Todos os outros campos amarelo-claros (clientes, churn, mix, custos, marketing) "
    "são EXEMPLOS estimados para a planilha já sair funcionando — troque pelos seus dados reais.":
        "Planilha em branco: comece pela aba Premissas e preencha os campos amarelo-claros. "
        "Os números aparecem sozinhos nas outras abas.",
    "Campos AMARELO-CLAROS = você preenche. Preços dos planos vieram do escopo; o restante são exemplos para substituir pelos seus números.":
        "Campos AMARELO-CLAROS = você preenche. Clique em cada um para ver uma dica.",
    "Campos AMARELO-CLAROS = você preenche. Preços vieram do escopo; o resto são exemplos para trocar pelos seus números.":
        "Campos AMARELO-CLAROS = você preenche. Clique em cada um para ver uma dica.",
    "A linha do mês 1 é um EXEMPLO de preenchimento — apague e lance os seus números.":
        "Lance uma linha por mês, sempre no fim do mês.",
    "Exemplo: contratar a partir do mês 7": "Ex.: se for contratar mais tarde, coloque o mês de início",
    "PREÇOS DOS PLANOS  (definidos no escopo)": "PREÇOS DOS PLANOS",
    "PREÇOS  (do escopo)": "PREÇOS",
}


def eh_campo(c):
    return c.fill.fill_type == "solid" and (c.fill.fgColor.rgb or "")[-6:] == INPUT


for origem, destino in ARQUIVOS.items():
    wb = load_workbook(origem)
    apagados = 0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in TEXTOS:
                    c.value = TEXTOS[c.value]
                elif eh_campo(c) and c.value is not None and not isinstance(c.value, str):
                    c.value = None      # números e datas; textos (nomes dos custos, cenário) ficam
                    apagados += 1
    ws = wb["Premissas"]
    for ref in ("C7", "C6"):   # mês de início (completa: C7, enxuta: C6)
        if ws[ref].number_format == "mmm/yy":
            ws[ref].value = "=DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)"
    wb.save(destino)
    print(f"ok -> {destino}  ({apagados} campos esvaziados)")
