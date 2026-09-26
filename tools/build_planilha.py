from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
import datetime as dt, sys
CLEAN = len(sys.argv) > 1 and sys.argv[1] == 'limpa'
OUT = sys.argv[2]

MUS='FFC21A'; INK='18171C'; PAPER='F6F5F2'; INPUT='FFF8E1'; LINE='D9D7DC'; MUTED='5E5C66'; OKF='E3F7EC'
F=lambda **k: Font(name='Arial', **{'size':10,'color':INK, **k})
fill=lambda c: PatternFill('solid', fgColor=c)
thin=Side(style='thin', color=LINE)
box=Border(bottom=thin)
BRL='"R$" #,##0.00;-"R$" #,##0.00;"-"'
PCT='0.0%;-0.0%;"-"'
NP=40; NV=1000; ND=200
P0=5; P1=P0+NP-1          # Produtos 5..44
V0=5; V1=V0+NV-1          # Vendas 5..1004
D0=13; D1=D0+ND-1         # Despesas 13..212

wb=Workbook()
def title(ws, t, sub, width_to='H'):
    ws.sheet_view.showGridLines=False
    ws['A1']=t; ws['A1'].font=F(size=16,bold=True)
    ws['A2']=sub; ws['A2'].font=F(color=MUTED)
    ws.row_dimensions[1].height=26
    for c in range(1, 12): ws.cell(row=3,column=c).fill=fill(MUS)
    ws.row_dimensions[3].height=4
def header(ws, row, labels, col=1):
    for i,l in enumerate(labels):
        c=ws.cell(row=row,column=col+i,value=l)
        c.font=F(bold=True,color='FFFFFF'); c.fill=fill(INK)
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=30
def inp(c): c.fill=fill(INPUT); c.font=F(color='0000FF')

# ---------- Como usar ----------
ws=wb.active; ws.title='Como usar'
title(ws,'Controle de Vendas da Barraca','Planilha grátis do Sai aê · pedido, cozinha e senha para feira, food truck e lanchonete')
rows=[
 ('COMO USAR',None),
 ('1. Produtos','Cadastre o que você vende, com o preço de venda e quanto custa para fazer cada um.'),
 ('2. Despesas','Confira as taxas da sua maquininha e lance as despesas do mês (taxa da feira, gás, embalagem...).'),
 ('3. Vendas','A cada venda, lance a data, o produto, a quantidade e a forma de pagamento. O resto a planilha calcula.'),
 ('4. Resumo','Escolha o mês e veja quanto vendeu, quanto sobrou de lucro e o que mais vende.'),
 ('',None),
 ('CORES',None),
 ('Amarelo-claro','Você preenche.'),
 ('Branco','A planilha calcula sozinha. Não precisa mexer.'),
 ('',None),
 ('MÊS DO RESUMO',None),
 ('Resumo','Abre no mês atual. Para ver outro mês, digite qualquer data dele na célula amarelo-clara.') if CLEAN else ('Exemplos','Esta versão vem com produtos, vendas e uma despesa de exemplo (setembro/2026) para você ver como funciona.'),
 ('',None),
 ('Cansou de anotar?','No Sai aê, cada venda lançada no caixa já entra no relatório do dia, sem digitar de novo.'),
]
r=5
for a,b in rows:
    ws.cell(row=r,column=1,value=a); ws.cell(row=r,column=2,value=b)
    if b is None and a: ws.cell(row=r,column=1).font=F(bold=True,color=MUTED,size=9)
    else: ws.cell(row=r,column=1).font=F(bold=True); ws.cell(row=r,column=2).font=F()
    r+=1
ws['A12'].fill=fill(INPUT); ws['A13'].border=Border(left=thin,right=thin,top=thin,bottom=thin)
ws['A18'].fill=fill(MUS); ws['B18'].fill=fill(MUS)
ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=100

# ---------- Produtos ----------
pr=wb.create_sheet('Produtos')
title(pr,'Produtos','Cadastre o que você vende. Preço e custo por unidade.')
header(pr,4,['Produto','Categoria','Preço de venda (R$)','Custo por unidade (R$)','Lucro por unidade (R$)','Margem'])
ex=[] if CLEAN else [('Pastel de carne','Pastéis',12,4.2),('Pastel de queijo','Pastéis',11,3.6),('Coxinha','Salgados',7,2.3),
    ('Caldo de cana 500ml','Bebidas',8,2.0),('Refrigerante lata','Bebidas',6,3.2)]
for i in range(NP):
    r=P0+i
    for c in 'ABCD': inp(pr[f'{c}{r}'])
    if i<len(ex):
        for c,v in zip('ABCD',ex[i]): pr[f'{c}{r}']=v
    pr[f'E{r}']=f'=IF(A{r}="","",C{r}-D{r})'
    pr[f'F{r}']=f'=IF(OR(A{r}="",C{r}=0),"",E{r}/C{r})'
    for c in 'CDE': pr[f'{c}{r}'].number_format=BRL
    pr[f'F{r}'].number_format=PCT
    for c in 'ABCDEF':
        pr[f'{c}{r}'].border=box
        if c in 'EF': pr[f'{c}{r}'].font=F()
pr['D4'].comment=Comment('Quanto custa para fazer 1 unidade: massa, recheio, óleo, embalagem.','Sai aê')
for c,w in zip('ABCDEF',[28,16,16,18,18,10]): pr.column_dimensions[c].width=w
pr.freeze_panes='A5'

# ---------- Despesas ----------
de=wb.create_sheet('Despesas')
title(de,'Despesas e taxas','Taxas da maquininha por forma de pagamento e as despesas do mês.')
header(de,4,['Forma de pagamento','Taxa (%)'])
taxas=[('Pix',0.0),('Dinheiro',0.0),('Débito',0.0199),('Crédito',0.0399)]
for i,(a,b) in enumerate(taxas):
    r=5+i; de[f'A{r}']=a; de[f'A{r}'].font=F(bold=True); de[f'B{r}']=b; inp(de[f'B{r}']); de[f'B{r}'].number_format='0.00%'
    de[f'A{r}'].border=box; de[f'B{r}'].border=box
de['C5']='Taxas de exemplo. Confira as da sua maquininha no app dela.'; de['C5'].font=F(color=MUTED,italic=True)
de['A11']='DESPESAS DO MÊS'; de['A11'].font=F(bold=True,color=MUTED,size=9)
header(de,12,['Data','Descrição','Categoria','Valor (R$)'])
exd=[] if CLEAN else [(dt.date(2026,9,5),'Saquinhos e guardanapos','Embalagem',48)]
for i in range(ND):
    r=D0+i
    for c in 'ABCD': inp(de[f'{c}{r}']); de[f'{c}{r}'].border=box
    if i<len(exd):
        for c,v in zip('ABCD',exd[i]): de[f'{c}{r}']=v
    de[f'A{r}'].number_format='DD/MM/YYYY'; de[f'D{r}'].number_format=BRL
dvc=DataValidation(type='list',formula1='"Ponto / taxa da feira,Gás,Embalagem,Ingredientes extras,Transporte,Ajudante,Outros"',allow_blank=True)
de.add_data_validation(dvc); dvc.add(f'C{D0}:C{D1}')
for c,w in zip('ABCD',[20,34,24,16]): de.column_dimensions[c].width=w
de.column_dimensions['C'].width=24

# ---------- Vendas ----------
ve=wb.create_sheet('Vendas')
title(ve,'Vendas','Uma linha por venda (ou por produto da venda). Preencha só as colunas amarelo-claras.')
header(ve,4,['Data','Produto','Qtd','Forma de pagamento','Preço (R$)','Total (R$)','Custo (R$)','Taxa (R$)','Lucro (R$)'])
exv=[] if CLEAN else [(dt.date(2026,9,5),'Pastel de carne',2,'Pix'),(dt.date(2026,9,5),'Caldo de cana 500ml',2,'Pix'),
     (dt.date(2026,9,5),'Coxinha',3,'Dinheiro'),(dt.date(2026,9,5),'Pastel de queijo',1,'Crédito'),
     (dt.date(2026,9,12),'Pastel de carne',4,'Débito'),(dt.date(2026,9,12),'Refrigerante lata',2,'Pix'),
     (dt.date(2026,9,12),'Pastel de queijo',2,'Pix'),(dt.date(2026,9,19),'Pastel de carne',3,'Crédito')]
PA=f'Produtos!$A${P0}:$A${P1}'
for i in range(NV):
    r=V0+i
    for c in 'ABCD': inp(ve[f'{c}{r}'])
    if i<len(exv):
        for c,v in zip('ABCD',exv[i]): ve[f'{c}{r}']=v
    ve[f'E{r}']=f'=IF(B{r}="","",IFERROR(INDEX(Produtos!$C${P0}:$C${P1},MATCH(B{r},{PA},0)),0))'
    ve[f'F{r}']=f'=IF(B{r}="","",C{r}*E{r})'
    ve[f'G{r}']=f'=IF(B{r}="","",C{r}*IFERROR(INDEX(Produtos!$D${P0}:$D${P1},MATCH(B{r},{PA},0)),0))'
    ve[f'H{r}']=f'=IF(B{r}="","",F{r}*IFERROR(INDEX(Despesas!$B$5:$B$8,MATCH(D{r},Despesas!$A$5:$A$8,0)),0))'
    ve[f'I{r}']=f'=IF(B{r}="","",F{r}-G{r}-H{r})'
    ve[f'A{r}'].number_format='DD/MM/YYYY'
    for c in 'EFGHI': ve[f'{c}{r}'].number_format=BRL; ve[f'{c}{r}'].font=F()
    for c in 'ABCDEFGHI': ve[f'{c}{r}'].border=box
dvp=DataValidation(type='list',formula1=f'={PA}',allow_blank=True,showErrorMessage=True,error='Escolha um produto da aba Produtos.')
dvf=DataValidation(type='list',formula1='"Pix,Dinheiro,Débito,Crédito"',allow_blank=True)
ve.add_data_validation(dvp); ve.add_data_validation(dvf)
dvp.add(f'B{V0}:B{V1}'); dvf.add(f'D{V0}:D{V1}')
for c,w in zip('ABCDEFGHI',[13,28,7,18,12,13,13,12,13]): ve.column_dimensions[c].width=w
ve.freeze_panes='A5'

# ---------- Resumo ----------
rs=wb.create_sheet('Resumo',1)
title(rs,'Resumo do mês','Escolha o mês na célula amarelo-clara. Todo o resto é automático.')
rs['A5']='Mês'; rs['A5'].font=F(bold=True)
rs['B5']='=TODAY()' if CLEAN else dt.date(2026,9,1); inp(rs['B5']); rs['B5'].number_format='MM/YYYY'; rs['B5'].font=F(color='0000FF',bold=True,size=12)
rs['B5'].comment=Comment('Digite qualquer dia do mês que você quer ver, ex.: 01/10/2026.','Sai aê')
rs['C5']='=DATE(YEAR(B5),MONTH(B5),1)'; rs['D5']='=EDATE(C5,1)'
for c in 'CD': rs[f'{c}5'].number_format='DD/MM/YYYY'; rs[f'{c}5'].font=F(color='B5B3BB',size=8)
rs['E5']='← início e fim do mês (automático)'; rs['E5'].font=F(color='B5B3BB',size=8)
VA=f'Vendas!$A${V0}:$A${V1}'
per=f'{VA},">="&$C$5,{VA},"<"&$D$5'
S=lambda col: f'SUMIFS(Vendas!${col}${V0}:${col}${V1},{per})'
header(rs,7,['Resultado do mês','Valor'])
res=[('Vendas',f'={S("F")}',BRL,False),
     ('(−) Custo dos produtos',f'={S("G")}',BRL,False),
     ('(−) Taxas da maquininha',f'={S("H")}',BRL,False),
     ('Lucro bruto','=B8-B9-B10',BRL,True),
     ('(−) Despesas do mês',f'=SUMIFS(Despesas!$D${D0}:$D${D1},Despesas!$A${D0}:$A${D1},">="&$C$5,Despesas!$A${D0}:$A${D1},"<"&$D$5)',BRL,False),
     ('Lucro líquido do mês','=B11-B12',BRL,True),
     ('Margem líquida','=IF(B8=0,0,B13/B8)',PCT,False),
     ('Itens vendidos',f'={S("C")}','#,##0;-#,##0;"-"',False),
     ('Dias com venda','=COUNTIF($J$8:$J$38,">0")','0;-0;"-"',False),
     ('Média de vendas por dia','=IF(B16=0,0,B8/B16)',BRL,False)]
for i,(a,f,nf,bold) in enumerate(res):
    r=8+i; rs[f'A{r}']=a; rs[f'B{r}']=f; rs[f'B{r}'].number_format=nf
    rs[f'A{r}'].font=F(bold=bold); rs[f'B{r}'].font=F(bold=bold)
    rs[f'A{r}'].border=box; rs[f'B{r}'].border=box
for c in 'AB': rs[f'{c}13'].fill=fill(MUS); rs[f'{c}13'].font=F(bold=True,size=12)

header(rs,7,['Forma de pagamento','Vendas','% do mês'],col=4)
for i,fp in enumerate(['Pix','Dinheiro','Débito','Crédito']):
    r=8+i; rs[f'D{r}']=fp; rs[f'D{r}'].font=F(bold=True)
    rs[f'E{r}']=f'=SUMIFS(Vendas!$F${V0}:$F${V1},Vendas!$D${V0}:$D${V1},D{r},{per})'
    rs[f'F{r}']=f'=IF($B$8=0,0,E{r}/$B$8)'
    rs[f'E{r}'].number_format=BRL; rs[f'F{r}'].number_format=PCT
    for c in 'DEF': rs[f'{c}{r}'].border=box

# por dia (para o gráfico)
header(rs,7,['Dia','Vendas'],col=9)
for k in range(31):
    r=8+k
    rs[f'I{r}']=f'=IF(MONTH($C$5+{k})=MONTH($C$5),$C$5+{k},"")'
    rs[f'J{r}']=f'=IF(I{r}="",0,SUMIFS(Vendas!$F${V0}:$F${V1},{VA},I{r}))'
    rs[f'I{r}'].number_format='DD/MM'; rs[f'J{r}'].number_format=BRL
    rs[f'I{r}'].font=F(color=MUTED); rs[f'J{r}'].font=F(color=MUTED)
ch=BarChart(); ch.type='col'; ch.title='Vendas por dia'; ch.style=10
ch.add_data(Reference(rs,min_col=10,min_row=7,max_row=38),titles_from_data=True)
ch.set_categories(Reference(rs,min_col=9,min_row=8,max_row=38))
ch.legend=None; ch.height=7.5; ch.width=17
ch.series[0].graphicalProperties.solidFill=MUS; ch.series[0].graphicalProperties.line.solidFill=MUS
ch.y_axis.numFmt='"R$" #,##0'; ch.y_axis.majorGridlines=None
rs.add_chart(ch,'D13')

# por produto
R0=20
rs[f'A{R0-1}']='POR PRODUTO NO MÊS'; rs[f'A{R0-1}'].font=F(bold=True,color=MUTED,size=9)
header(rs,R0,['Produto','Qtd vendida','Vendas','Lucro','% das vendas','Lucro por unidade'])
for i in range(NP):
    r=R0+1+i; pr_r=P0+i
    rs[f'A{r}']=f'=IF(Produtos!A{pr_r}="","",Produtos!A{pr_r})'
    crit=f'Vendas!$B${V0}:$B${V1},A{r},{per}'
    rs[f'B{r}']=f'=IF(A{r}="","",SUMIFS(Vendas!$C${V0}:$C${V1},{crit}))'
    rs[f'C{r}']=f'=IF(A{r}="","",SUMIFS(Vendas!$F${V0}:$F${V1},{crit}))'
    rs[f'D{r}']=f'=IF(A{r}="","",SUMIFS(Vendas!$I${V0}:$I${V1},{crit}))'
    rs[f'E{r}']=f'=IF(OR(A{r}="",$B$8=0),"",C{r}/$B$8)'
    rs[f'F{r}']=f'=IF(A{r}="","",Produtos!E{pr_r})'
    rs[f'B{r}'].number_format='#,##0;-#,##0;"-"'
    for c in 'CDF': rs[f'{c}{r}'].number_format=BRL
    rs[f'E{r}'].number_format=PCT
    for c in 'ABCDEF': rs[f'{c}{r}'].border=box
for c,w in zip('ABCDEFGHIJ',[28,16,16,16,14,16,2,2,10,13]): rs.column_dimensions[c].width=w
rs.column_dimensions['G'].width=4; rs.column_dimensions['H'].width=4

for s in wb.worksheets:
    s.sheet_properties.tabColor = MUS if s.title in ('Resumo','Vendas') else INK
wb.active=1
wb.save(OUT)
