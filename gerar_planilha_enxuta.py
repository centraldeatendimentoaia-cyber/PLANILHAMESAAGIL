# -*- coding: utf-8 -*-
"""Gera a versão ENXUTA da planilha do Mesa Ágil (Mesa_Agil_Enxuta.xlsx).

Para a rotina do dia a dia: 4 abas, ~15 campos para preencher uma vez e 6 números por mês.
A versão completa (gerar_planilha.py) continua para apresentar a sócio/investidor.
Mesmo visual: estilo_app.py (design system do app).
Rode:  python3 gerar_planilha_enxuta.py
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from estilo_app import *  # noqa: F401,F403  (cores, fontes, put, entrada, botao, tile, card...)

SAIDA = "Mesa_Agil_Enxuta.xlsx"
PAINEL, PR, PJ, MM = "Painel", "Premissas", "Projeção", "Meu Mês"
NAV = dict(inicio=("⌂  Painel", PAINEL), painel=None)
FR, LR = 6, 29  # meses 1..24

wb = Workbook()
wb.remove(wb.active)

# ====================================================================== PREMISSAS
pr = base(wb, PR, "PREMISSAS  ·  PREENCHA UMA VEZ", "G", AMBAR_CTA,
          "Campos VERDE-CLAROS = você preenche. Preços vieram do escopo; o resto são exemplos para trocar pelos seus números.",
          {"B": 46, "C": 15, "D": 15, "E": 15, "F": 15, "G": 56}, **NAV)

secao(pr, "B5", "GERAL")
put(pr, "B6", "Mês de início"); entrada(pr, "C6", dt.date(2026, 10, 1), fmt=MES)
put(pr, "D6", "Primeiro mês de vendas", 9, cor=GRAFITE)
put(pr, "B7", "Caixa inicial (R$)"); entrada(pr, "C7", 0, fmt=MOEDA)
put(pr, "D7", "Dinheiro que já existe para bancar o começo", 9, cor=GRAFITE)
put(pr, "B8", "Cenário em uso", b=True); entrada(pr, "C8", "Realista", al="center")
put(pr, "D8", "Pessimista, Realista ou Otimista", 9, cor=GRAFITE)
dv = DataValidation(type="list", formula1='"Pessimista,Realista,Otimista"', allow_blank=False,
                    showInputMessage=True, promptTitle="Cenário", prompt="Escolha o cenário usado na Projeção e no Painel.")
pr.add_data_validation(dv)
dv.add("C8")

secao(pr, "B10", "PREÇOS  (do escopo)")
cabecalho(pr, 11, 2, ["PLANO", "MENSAL", "DESCONTO ANUAL", "ANUAL (POR MÊS)", "ANUAL (COBRADO)"], altura=26)
for k, (nome, preco, desc) in enumerate([("Essencial", 57.90, 0.15), ("Pro", 87.90, 0.20)]):
    r = 12 + k
    put(pr, f"B{r}", nome, b=True, borda=borda_linha)
    entrada(pr, f"C{r}", preco, fmt=MOEDA)
    entrada(pr, f"D{r}", desc, fmt=PCT)
    put(pr, f"E{r}", f"=C{r}*(1-D{r})", fmt=MOEDA, borda=borda_linha)
    put(pr, f"F{r}", f"=ROUND(E{r}*12,2)", fmt=MOEDA, borda=borda_linha)

secao(pr, "B15", "CLIENTES  ·  POR CENÁRIO")
cabecalho(pr, 16, 2, ["PREMISSA", "PESSIMISTA", "REALISTA", "OTIMISTA", "EM USO", "O QUE É"], altura=26)
cen = [
    ("Novos clientes no mês 1", (4, 8, 15), NUM1, "Assinaturas novas no primeiro mês"),
    ("Crescimento dos novos clientes (ao mês)", (0.03, 0.07, 0.12), PCT, "Ex.: 7% → 8; 8,6; 9,2 novos..."),
    ("Churn mensal", (0.08, 0.05, 0.03), PCT, "% dos clientes que cancelam por mês"),
    ("% no plano Pro", (0.15, 0.25, 0.35), PCT, "O restante fica no Essencial"),
    ("% que paga anual", (0.10, 0.20, 0.30), PCT, "O restante paga mensal"),
]
C = {}
for k, (rot, vals, fmt, oque) in enumerate(cen):
    r = 17 + k
    put(pr, f"B{r}", rot, borda=borda_linha)
    for j, v in enumerate(vals):
        entrada(pr, f"{L(3 + j)}{r}", v, fmt=fmt)
    put(pr, f"F{r}", f"=INDEX(C{r}:E{r},MATCH($C$8,$C$16:$E$16,0))", b=True, fmt=fmt, al="right", borda=borda_linha)
    put(pr, f"G{r}", oque, 9, cor=GRAFITE)
    C[k] = f"Premissas!$F${r}"
NOVOS1, CRESC, CHURN, PRO, ANUAL = C[0], C[1], C[2], C[3], C[4]

secao(pr, "B23", "CUSTOS")
put(pr, "B24", "Taxas sobre o faturamento (imposto + gateway)", borda=borda_linha)
entrada(pr, "C24", 0.095, fmt=PCT)
put(pr, "D24", "Ex.: Simples 6% + gateway 3,5%. Confirme com o contador.", 9, cor=GRAFITE)
put(pr, "B25", "Anúncios por mês (R$)", borda=borda_linha)
entrada(pr, "C25", 1500, fmt=MOEDA)
put(pr, "D25", "Meta/Google/Instagram + comissão de indicação", 9, cor=GRAFITE)

secao(pr, "B27", "CUSTOS FIXOS DO MÊS")
cabecalho(pr, 28, 2, ["DESCRIÇÃO", "VALOR MENSAL"], altura=24)
fixos = [
    ("Hospedagem, banco de dados e tempo real", 150),
    ("Domínio e e-mail", 40),
    ("Ferramentas (IA, design, site)", 205),
    ("WhatsApp de atendimento", 50),
    ("Pró-labore", 1500),
    ("Contador", 250),
    ("", None),
    ("", None),
]
FX1, FX2 = 29, 29 + len(fixos) - 1
for k, (d, v) in enumerate(fixos):
    entrada(pr, f"B{FX1 + k}", d, al="left")
    entrada(pr, f"C{FX1 + k}", v, fmt=MOEDA)
TOT = FX2 + 1
put(pr, f"B{TOT}", "TOTAL DE CUSTOS FIXOS", b=True, borda=borda_total)
put(pr, f"C{TOT}", f"=SUM(C{FX1}:C{FX2})", b=True, fmt=MOEDA, borda=borda_total)

secao(pr, f"B{TOT + 2}", "CALCULADO  (não precisa mexer)")
calc = [
    ("Ticket médio por cliente/mês",
     f"=(1-{PRO})*((1-{ANUAL})*C12+{ANUAL}*F12/12)+{PRO}*((1-{ANUAL})*C13+{ANUAL}*F13/12)", MOEDA,
     "Mistura de planos e de mensal/anual"),
    ("Sobra por cliente/mês (depois das taxas)", f"=C{TOT + 3}*(1-C24)", MOEDA, "É o que cada cliente paga das contas"),
    ("Valor médio pago à vista no plano anual", f"=(1-{PRO})*F12+{PRO}*F13", MOEDA, "Entra de uma vez no caixa"),
    ("Contas do mês (fixos + anúncios)", f"=C{TOT}+C25", MOEDA, ""),
]
for k, (rot, f_, fmt, obs) in enumerate(calc):
    r = TOT + 3 + k
    put(pr, f"B{r}", rot, borda=borda_linha)
    put(pr, f"C{r}", f_, b=True, fmt=fmt, borda=borda_linha)
    put(pr, f"D{r}", obs, 9, cor=GRAFITE)
TICKET, SOBRA, ANUAL_VISTA, CONTAS = (f"Premissas!$C${TOT + 3 + k}" for k in range(4))

for ref, t, txt in [
    ("C6", "Início", "Data do primeiro mês de vendas. Ex.: 01/10/2026."),
    ("C7", "Caixa inicial", "Quanto dinheiro já existe (R$). Pode ser 0."),
    ("C17:E17", "Novos no mês 1", "Quantas assinaturas novas no primeiro mês."),
    ("C18:E18", "Crescimento", "Quanto os novos clientes crescem a cada mês. Digite 7% ou 0,07."),
    ("C19:E19", "Churn", "% dos clientes que cancelam por mês. Digite 5% ou 0,05."),
    ("C20:E20", "Plano Pro", "% dos clientes no Pro. Digite 25% ou 0,25."),
    ("C21:E21", "Plano anual", "% dos clientes que pagam anual. Digite 20% ou 0,20."),
    ("C24", "Taxas", "Imposto (Simples) + taxa do gateway, somados. Ex.: 9,5%."),
    ("C25", "Anúncios", "Quanto vai gastar por mês para conquistar clientes (R$)."),
    (f"C{FX1}:C{FX2}", "Custo fixo", "Valor por mês (R$). Custo anual? Divida por 12."),
]:
    dica(pr, ref, t, txt)
pr.freeze_panes = "A4"

# ====================================================================== PROJEÇÃO
pj = base(wb, PJ, "", "L", GRAFITE,
          "Mês a mês no cenário em uso. O plano anual entra inteiro no caixa quando é vendido ou renovado. Tudo automático.",
          {"B": 7, "C": 10, **{L(c): 15 for c in range(4, 12)}, "L": 12}, **NAV)
pj["B2"].value = '="PROJEÇÃO 24 MESES  ·  CENÁRIO "&UPPER(Premissas!$C$8)'
cabecalho(pj, 5, 2, ["Nº", "MÊS", "NOVOS", "CANCELADOS", "CLIENTES ATIVOS", "MRR", "CUSTOS", "LUCRO",
                     "CAIXA DO MÊS", "CAIXA ACUMULADO", "MÊS COM LUCRO?"], altura=32)
pj["L5"].font = font(8, True, GRAFITE)
for r in range(FR, LR + 1):
    f = {
        "B": "=1" if r == FR else f"=B{r - 1}+1",
        "C": f"=EDATE(Premissas!$C$6,B{r}-1)",
        "D": f"={NOVOS1}*(1+{CRESC})^(B{r}-1)",
        "E": f"=N(F{r - 1})*{CHURN}",
        "F": f"=N(F{r - 1})-E{r}+D{r}",
        "G": f"=F{r}*{TICKET}",
        "H": f"=G{r}*Premissas!$C$24+{CONTAS}",
        "I": f"=G{r}-H{r}",
        # caixa = lucro + anual recebido à vista (novos + renovações) − parte do anual que já é receita do mês
        "J": (f"=I{r}+(D{r}+IF(B{r}>12,INDEX($D${FR}:$D${LR},B{r}-12)*(1-{CHURN})^12,0))*{ANUAL}*{ANUAL_VISTA}"
              f"-F{r}*{ANUAL}*{ANUAL_VISTA}/12"),
        "K": f"=IF(B{r}=1,Premissas!$C$7,N(K{r - 1}))+J{r}",
        "L": f'=IF(I{r}>0,B{r},"")',
    }
    fmts = {"B": "0", "C": MES, "D": NUM1, "E": NUM1, "F": NUM1, "G": MOEDA, "H": MOEDA, "I": MOEDA,
            "J": MOEDA, "K": MOEDA, "L": "0"}
    for col, frm in f.items():
        put(pj, f"{col}{r}", frm, fmt=fmts[col], borda=borda_linha, b=col in ("F", "I", "K"),
            al="center" if col in ("B", "C", "L") else None, cor=GRAFITE if col == "L" else PRETO)
T = LR + 1
put(pj, f"B{T}", "TOTAL", b=True, borda=borda_total)
pj.merge_cells(f"B{T}:C{T}")
for col in "DEGHIJ":
    put(pj, f"{col}{T}", f"=SUM({col}{FR}:{col}{LR})", b=True, fmt=NUM1 if col in "DE" else MOEDA, borda=borda_total)
for col in "FKL":
    put(pj, f"{col}{T}", None, borda=borda_total)
for col in "IJK":
    cor_condicional(pj, f"{col}{FR}:{col}{T}", f"{col}{FR}<0")
pj.freeze_panes = "D6"

# ====================================================================== MEU MÊS
mm = base(wb, MM, "MEU MÊS  ·  O QUE ACONTECEU DE VERDADE", "N", GRAFITE,
          "No fim de cada mês, preencha os 6 campos verde-claros da linha do mês. O resto compara com a projeção sozinho.",
          {"B": 7, "C": 10, **{L(c): 14 for c in range(4, 15)}}, **NAV)
mm.merge_cells("D4:I4")
put(mm, "D4", "VOCÊ LANÇA", 9, True, "FFFFFF", bg=MARCA_2, al="center")
mm.merge_cells("J4:N4")
put(mm, "J4", "A PLANILHA CALCULA", 9, True, "FFFFFF", bg=MARCA_2, al="center")
cabecalho(mm, 5, 2, ["Nº", "MÊS", "CLIENTES ATIVOS (FIM DO MÊS)", "NOVOS", "CANCELADOS", "RECEBIDO (R$)",
                     "GASTO TOTAL (R$)", "DESSE GASTO, ANÚNCIOS (R$)",
                     "CLIENTES PROJETADOS", "REAL × PROJETADO", "CHURN REAL", "CAC REAL", "SOBROU NO CAIXA"], altura=42)
exemplo = {"D": 8, "E": 8, "F": 0, "G": 520, "H": 3100, "I": 1400}
for r in range(FR, LR + 1):
    put(mm, f"B{r}", f"='{PJ}'!B{r}", fmt="0", al="center", borda=borda_linha)
    put(mm, f"C{r}", f"='{PJ}'!C{r}", fmt=MES, al="center", borda=borda_linha)
    for col, fmt in zip("DEFGHI", [NUM, NUM, NUM, MOEDA, MOEDA, MOEDA]):
        entrada(mm, f"{col}{r}", exemplo[col] if r == FR else None, fmt=fmt)
    put(mm, f"J{r}", f"='{PJ}'!F{r}", fmt=NUM, borda=borda_linha, cor=GRAFITE)
    put(mm, f"K{r}", f'=IF(D{r}="","",IFERROR(D{r}/J{r}-1,""))', fmt='+0%;-0%;0%', borda=borda_linha, b=True)
    put(mm, f"L{r}", f'=IF(OR(D{r}="",F{r}=""),"",IFERROR(F{r}/(D{r}-E{r}+F{r}),""))', fmt=PCT, borda=borda_linha)
    put(mm, f"M{r}", f'=IF(OR(I{r}="",E{r}=""),"",IFERROR(I{r}/E{r},""))', fmt=MOEDA, borda=borda_linha)
    put(mm, f"N{r}", f'=IF(OR(G{r}="",H{r}=""),"",G{r}-H{r})', fmt=MOEDA, borda=borda_linha, b=True)
cor_condicional(mm, f"K{FR}:K{LR}", f"AND(ISNUMBER(K{FR}),K{FR}<-0.1)", f"AND(ISNUMBER(K{FR}),K{FR}>=0)")
cor_condicional(mm, f"N{FR}:N{LR}", f"AND(ISNUMBER(N{FR}),N{FR}<0)", f"AND(ISNUMBER(N{FR}),N{FR}>0)")
put(mm, f"B{LR + 2}", "A linha do mês 1 é um EXEMPLO de preenchimento — apague e lance os seus números.", 9, True, AMBAR_CTA)
put(mm, f"B{LR + 3}", "Churn real = cancelados ÷ clientes no início do mês.  CAC real = anúncios ÷ novos clientes.  "
    "Real × projetado em vermelho = mais de 10% abaixo do planejado.", 9, cor=GRAFITE)
for ref, t, txt in [
    (f"D{FR}:D{LR}", "Clientes ativos", "Quantas assinaturas ativas no último dia do mês."),
    (f"E{FR}:E{LR}", "Novos", "Quantos clientes novos pagaram neste mês."),
    (f"F{FR}:F{LR}", "Cancelados", "Quantos cancelaram neste mês."),
    (f"G{FR}:G{LR}", "Recebido", "Tudo que entrou na conta no mês (mensalidades + anuais), já sem a taxa do gateway."),
    (f"H{FR}:H{LR}", "Gasto total", "Tudo que saiu no mês: custos fixos, anúncios, impostos, pró-labore."),
    (f"I{FR}:I{LR}", "Anúncios", "Quanto do gasto total foi com anúncios/indicação (para calcular o CAC)."),
]:
    dica(mm, ref, t, txt)
mm.freeze_panes = "D6"

# ====================================================================== PAINEL
TILES = [2, 5, 8, 11, 14, 17]
db = base(wb, PAINEL, "", "R", MARCA, larguras={L(c): (3 if (c - 1) % 3 == 0 else 12) for c in range(2, 19)}, **NAV)
db["B2"].value = '="MESA ÁGIL  ·  PAINEL  ·  CENÁRIO "&UPPER(Premissas!$C$8)'
db["B2"].font = font(16, True, "FFFFFF")
botao(db, "B1", "⚙  Premissas", f"#'{PR}'!A1", CARD, PRETO, mescla="B1:C1", borda=CLARO)
botao(db, "E1", "✎  Meu Mês", f"#'{MM}'!A1", CARD, PRETO, mescla="E1:F1", borda=CLARO)
botao(db, "H1", "↗  Projeção", f"#'{PJ}'!A1", CARD, PRETO, mescla="H1:I1", borda=CLARO)
put(db, "B3", "Como usar:  1) Premissas — ajuste uma vez.   2) Meu Mês — no fim de cada mês, lance 6 números.   3) Painel — veja se está no rumo.",
    9, cor=GRAFITE)
botao(db, "Q3", "Lançar o mês  →", f"#'{MM}'!D6", AMBAR_CTA, "FFFFFF", mescla="Q3:R3")

P_ = lambda col, r: f"'{PJ}'!${col}${r}"
M_ = lambda col: f"'{MM}'!${col}${FR}:${col}${LR}"
N_MESES = f"COUNT({M_('D')})"

secao(db, "B5", "O PLANO  ·  O NEGÓCIO SE PAGA?")
tile(db, TILES[0], 6, "EQUILÍBRIO", f"=IFERROR(ROUNDUP({CONTAS}/{SOBRA},0),0)", '#,##0" clientes"', "pagam fixos + anúncios")
cap = tile(db, TILES[1], 6, "CAPITAL NECESSÁRIO", f"=MAX(0,-MIN({P_('K', FR)}:{P_('K', LR)}))", MOEDA0, "maior buraco no caixa")
tile(db, TILES[2], 6, "1º MÊS COM LUCRO",
     f"=IF(COUNT({P_('L', FR)}:{P_('L', LR)})=0,\"Não atinge\",INDEX({P_('C', FR)}:{P_('C', LR)},MIN({P_('L', FR)}:{P_('L', LR)})))",
     MES, "no cenário em uso")
ltv = tile(db, TILES[3], 6, "LTV / CAC",
           f"=IFERROR(({SOBRA}/{CHURN})/(Premissas!$C$25*24/{P_('D', T)}),0)", VEZES,
           f"=IFERROR({SOBRA}/{CHURN},0)", '"LTV R$ "#,##0" · ideal ≥ 3x"')
tile(db, TILES[4], 6, "CLIENTES — MÊS 12", f"={P_('F', FR + 11)}", NUM, f"={P_('F', LR)}", '"mês 24: "#,##0')
tile(db, TILES[5], 6, "MRR — MÊS 12", f"={P_('G', FR + 11)}", MOEDA0, f"={P_('G', LR)}", '"mês 24: R$ "#,##0')
db.conditional_formatting.add(cap, FormulaRule(formula=[f"{cap}>0"], font=Font(color=VERMELHO)))
db.conditional_formatting.add(cap, FormulaRule(formula=[f"{cap}=0"], font=Font(color=VERDE)))
db.conditional_formatting.add(ltv, FormulaRule(formula=[f"{ltv}>=3"], font=Font(color=VERDE)))
db.conditional_formatting.add(ltv, FormulaRule(formula=[f"{ltv}<1"], font=Font(color=VERMELHO)))

secao(db, "B10", "A REALIDADE  ·  O QUE VOCÊ LANÇOU NO MEU MÊS")
tile(db, TILES[0], 11, "MESES LANÇADOS", f"={N_MESES}", '0" de 24"', "linhas preenchidas")
rp = tile(db, TILES[1], 11, "CLIENTES REAL × PLANO",
          f"=IF({N_MESES}=0,\"—\",IFERROR(INDEX({M_('D')},{N_MESES})/INDEX({M_('J')},{N_MESES})-1,\"—\"))",
          '+0%;-0%;0%', f"=IF({N_MESES}=0,0,INDEX({M_('D')},{N_MESES}))", '"hoje: "#,##0" clientes"')
tile(db, TILES[2], 11, "CHURN REAL", f"=IFERROR(AVERAGE({M_('L')}),\"—\")", PCT, f"={CHURN}", '"no plano: "0.0%')
tile(db, TILES[3], 11, "CAC REAL", f"=IFERROR(SUM({M_('I')})/SUM({M_('E')}),\"—\")", MOEDA0,
     f"=IFERROR(Premissas!$C$25*24/{P_('D', T)},0)", '"no plano: R$ "#,##0')
tile(db, TILES[4], 11, "SOBROU NO CAIXA", f"=SUM({M_('N')})", MOEDA0, "soma dos meses lançados", neg=True)
tile(db, TILES[5], 11, "CAIXA NO PLANO", f"=SUMPRODUCT(({P_('B', FR)}:{P_('B', LR)}<={N_MESES})*{P_('J', FR)}:{P_('J', LR)})",
     MOEDA0, "mesmos meses, projetado", neg=True)
db.conditional_formatting.add(rp, FormulaRule(formula=[f"AND(ISNUMBER({rp}),{rp}<-0.1)"], font=Font(color=VERMELHO)))
db.conditional_formatting.add(rp, FormulaRule(formula=[f"AND(ISNUMBER({rp}),{rp}>=0)"], font=Font(color=VERDE)))

secao(db, "B15", "EVOLUÇÃO")
cats = Reference(pj, min_col=3, min_row=FR, max_row=LR)
g1 = novo("line", "Clientes ativos: plano × real", "#,##0")
g1.add_data(Reference(pj, min_col=6, min_row=FR, max_row=LR))
estilo_serie(g1.series[-1], CINZA, "dash", 28000)
g1.series[-1].tx = SeriesLabel(v="Plano")
g1.add_data(Reference(mm, min_col=4, min_row=FR, max_row=LR))
estilo_serie(g1.series[-1], VERDE, None, 32000)
g1.series[-1].tx = SeriesLabel(v="Real")
g1.series[-1].marker.symbol = "circle"
g1.series[-1].marker.size = 6
g1.series[-1].marker.graphicalProperties = GraphicalProperties(solidFill=VERDE, ln=LineProperties(solidFill=VERDE))
g1.set_categories(cats)
db.add_chart(g1, "B16")
g2 = novo("line", "Caixa acumulado (plano)", '"R$ "#,##0')
g2.add_data(Reference(pj, min_col=11, min_row=5, max_row=LR), titles_from_data=True)
estilo_serie(g2.series[-1], VERDE, None, 28000)
g2.set_categories(cats)
g2.legend = None
db.add_chart(g2, "K16")

secao(db, "B33", "ALERTAS")
alertas = [
    (f"{cap}>0", "Você vai precisar de capital: o caixa fica negativo no começo (veja o card CAPITAL NECESSÁRIO)",
     "O caixa nunca fica negativo no plano"),
    (f"{ltv}<3", "LTV/CAC abaixo de 3x — cada cliente devolve pouco do que custa para conquistar",
     "LTV/CAC saudável: cada cliente devolve 3x ou mais o que custou"),
    (f"AND(ISNUMBER({rp}),{rp}<-0.1)", "Você está mais de 10% abaixo do plano em clientes — reveja anúncios ou churn",
     "Clientes dentro do plano (ou ainda sem lançamentos)"),
]
AUX = 21  # U (oculta)
for k, (cond, ruim, bom) in enumerate(alertas):
    r = 34 + k
    put(db, f"{L(AUX)}{r}", f"={cond}", 8, cor=GRAFITE)
    db.merge_cells(f"B{r}:R{r}")
    put(db, f"B{r}", f'=IF({L(AUX)}{r},"✖   {ruim}","✔   {bom}")', 11, True, al="left")
    db.row_dimensions[r].height = 26
    db.conditional_formatting.add(f"B{r}:R{r}", FormulaRule(
        formula=[f"${L(AUX)}${r}"], font=Font(color="991B1B", bold=True), fill=fill("FEF2F2"),
        border=Border(top=Side("thin", color="FCA5A5"), bottom=Side("thin", color="FCA5A5"))))
    db.conditional_formatting.add(f"B{r}:R{r}", FormulaRule(
        formula=[f"NOT(${L(AUX)}${r})"], font=Font(color="065F46", bold=True), fill=fill("ECFDF5"),
        border=Border(top=Side("thin", color="10B981"), bottom=Side("thin", color="10B981"))))
db.column_dimensions[L(AUX)].hidden = True
put(db, "B38", "Quer mais detalhe (plano anual separado, unit economics por plano, três cenários lado a lado)? "
    "Use a versão completa: Mesa_Agil_Financeiro.xlsx.", 9, cor=GRAFITE)

finalizar(wb, [PAINEL, PR, MM, PJ], PAINEL)
wb.save(SAIDA)
print("ok ->", SAIDA)
