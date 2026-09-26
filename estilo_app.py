# -*- coding: utf-8 -*-
"""Estilo compartilhado das planilhas do Saiaê — design system Saiaê v1.0.

Mostarda #FFC21A = cor principal (um destaque por tela; texto sobre mostarda é sempre tinta),
tinta #18171C = texto e fundos escuros, papel #F6F5F2 = fundo, branco = superfícies.
Outfit (títulos e números) + Figtree (interface). Status: no prazo / atenção / atrasado / info."""
from openpyxl.chart import BarChart, LineChart
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import MergedCell

# ---------------------------------------------------------------- tokens Saiaê v1.0
# tinta (neutros)
PRETO, GRAFITE, CINZA, CLARO, LINHA = "18171C", "5E5C66", "8A8893", "B5B3BB", "D9D7DC"  # 900/500/400/300/200
MARCA, MARCA_2 = "18171C", "2E2C35"            # barra de título (tinta 900) / faixas de grupo (tinta 700)
# mostarda
MOSTARDA, MOSTARDA_600, MOSTARDA_700, MOSTARDA_50 = "FFC21A", "E6A800", "B38300", "FFF8E1"
AMBAR_CTA = MOSTARDA                           # botão primário (um por tela), texto sempre tinta
# status semânticos: cor / fundo / texto
OK_COR, OK_BG, OK_TXT = "1FBF6A", "E3F7EC", "0E7A40"        # no prazo (positivo)
ATN_COR, ATN_BG, ATN_TXT = "FF8A1C", "FFF0E0", "A34E00"     # atenção
ERR_COR, ERR_BG, ERR_TXT = "F0412F", "FDE7E4", "B22514"     # atrasado / erro (negativo)
INFO_COR, INFO_BG, INFO_TXT = "2F7BF5", "E6F0FE", "1A4FA8"  # info
VERDE, VERMELHO, AMBAR = OK_TXT, ERR_TXT, ATN_COR           # texto positivo / negativo (contraste AA)
# superfícies
CANVAS, SUB, CARD = "F6F5F2", "ECEAE6", "FFFFFF"            # papel 50 / tinta 100 / branco
INPUT, INPUT_TXT, INPUT_BORDA = MOSTARDA_50, PRETO, MOSTARDA_600  # campo que você preenche
# gráficos
SERIE1, SERIE_BAR = MOSTARDA_600, MOSTARDA
# tipografia (gratuitas no Google Fonts)
F, F_DISPLAY = "Figtree", "Outfit"

MOEDA = '"R$ "#,##0.00;"-R$ "#,##0.00;"-"'
MOEDA0 = '"R$ "#,##0;"-R$ "#,##0;"R$ 0"'
PCT = '0.0%;-0.0%;"-"'
NUM = '#,##0;-#,##0;"-"'
NUM1 = '#,##0.0;-#,##0.0;"-"'
MES = 'mmm/yy'
VEZES = '0.0"x";-0.0"x";"-"'
MESES = '0.0" meses";-0.0" meses";"-"'

fill = lambda c: PatternFill("solid", fgColor=c)
_s = lambda c: Side("thin", color=c)
borda_input = Border(left=_s(INPUT_BORDA), right=_s(INPUT_BORDA), top=_s(INPUT_BORDA), bottom=_s(INPUT_BORDA))
borda_linha = Border(left=_s(LINHA), right=_s(LINHA), top=_s(LINHA), bottom=_s(LINHA))
borda_total = Border(left=_s(LINHA), right=_s(LINHA), top=Side("medium", color=CLARO), bottom=_s(LINHA))


def openpyxl_col(letra):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letra)


def font(sz=11, b=False, cor=PRETO, i=False, fam=None):
    return Font(name=fam or F, size=sz, bold=b, color=cor, italic=i)


def put(ws, ref, valor, sz=11, b=False, cor=PRETO, fmt=None, bg=None, al=None, borda=None, wrap=False, i=False, fam=None):
    c = ws[ref]
    c.value = valor
    c.font = font(sz, b, cor, i, fam)
    if fmt:
        c.number_format = fmt
    if bg:
        c.fill = fill(bg)
    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if borda:
        c.border = borda
        if borda is borda_total and not bg:
            c.fill = fill(SUB)
    return c


def entrada(ws, ref, valor, fmt=None, al="right"):
    """Campo que você preenche: fundo mostarda 50, borda mostarda 600, texto tinta."""
    return put(ws, ref, valor, cor=INPUT_TXT, b=True, fmt=fmt, bg=INPUT, al=al, borda=borda_input)


def botao(ws, ref, texto, link, bg, cor, mescla=None, borda=None):
    """Botão: primário (mostarda, texto tinta) ou secundário (branco com borda tinta 300)."""
    if mescla:
        ws.merge_cells(mescla)
    c = put(ws, ref, texto, 10, True, cor, bg=bg, al="center")
    c.border = Border(left=_s(borda or bg), right=_s(borda or bg), top=_s(borda or bg), bottom=_s(borda or bg))
    c.hyperlink = link
    return c


def dica(ws, ref, titulo, texto):
    dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle=titulo[:32], prompt=texto[:255])
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add(ref)


def base(wb, nome, titulo, ultima_col, aba_cor, subtitulo=None, larguras=None,
         inicio=("⌂  Início", "Comece Aqui"), painel=("▦  Dashboard", "Dashboard")):
    """Aba no padrão Saiaê: botões de navegação, barra de título tinta com filete mostarda e subtítulo."""
    ws = wb.create_sheet(nome)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = aba_cor
    ws.column_dimensions["A"].width = 3
    for col, w in (larguras or {}).items():
        ws.column_dimensions[col].width = w
    if nome != inicio[1]:
        botao(ws, "B1", inicio[0], f"#'{inicio[1]}'!A1", CARD, PRETO, borda=CLARO)
        if painel and nome != painel[1]:
            botao(ws, "C1", painel[0], f"#'{painel[1]}'!A1", CARD, PRETO, borda=CLARO)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(f"B2:{ultima_col}2")
    put(ws, "B2", titulo, 16, True, "FFFFFF", bg=MARCA, al="left", fam=F_DISPLAY)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in range(2, openpyxl_col(ultima_col) + 1):
        ws.cell(2, cc).border = Border(bottom=Side("thick", color=MOSTARDA))
    ws.row_dimensions[2].height = 42
    if subtitulo:
        put(ws, "B3", subtitulo, 9, cor=GRAFITE)
    return ws


def secao(ws, ref, texto):
    put(ws, ref, texto, 11, True, PRETO, fam=F_DISPLAY)


def cabecalho(ws, linha, col_ini, textos, altura=30):
    for k, t in enumerate(textos):
        put(ws, f"{L(col_ini + k)}{linha}", t, 9, True, PRETO, bg=SUB, al="center", wrap=True, borda=borda_linha)
    ws.row_dimensions[linha].height = altura


def cor_condicional(ws, rng, cond_neg, cond_pos=None):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[cond_neg], font=Font(color=VERMELHO)))
    if cond_pos:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[cond_pos], font=Font(color=VERDE)))


def estilo_serie(s, cor, tracejado=None, largura=22000):
    s.graphicalProperties.line.solidFill = cor
    s.graphicalProperties.line.width = largura
    if tracejado:
        s.graphicalProperties.line.dashStyle = tracejado
    s.smooth = False


def estilo_grafico(ch, titulo, fmt_y):
    """Card de gráfico: fundo branco, borda slate, sem grade pesada, fonte do app."""
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.line import LineProperties
    from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
    from openpyxl.chart.title import Title
    from openpyxl.chart.text import Text
    from openpyxl.drawing.text import RegularTextRun
    ct = CharacterProperties(latin=DFont(typeface=F_DISPLAY), sz=1200, b=True, solidFill=PRETO)
    ch.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=ct), r=[RegularTextRun(rPr=ct, t=titulo)])])),
                     overlay=False)
    ch.plot_visible_only = False
    ch.x_axis.tickLblPos = "low"
    ch.height, ch.width = 8, 16
    ch.legend.position = "b"
    ch.y_axis.number_format = fmt_y
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINHA))
    ch.x_axis.number_format = MES
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.graphical_properties = GraphicalProperties(solidFill=CARD, ln=LineProperties(solidFill=LINHA))
    cp = CharacterProperties(latin=DFont(typeface=F), sz=900, solidFill=GRAFITE)
    ch.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def card(ws, c1, c2, r1, r2, borda_cor=LINHA):
    """Card: superfície branca, borda tinta 200."""
    for rr in range(r1, r2 + 1):
        for cc in range(openpyxl_col(c1), openpyxl_col(c2) + 1):
            cel = ws.cell(rr, cc)
            lado = Side("thin", color=borda_cor)
            cel.border = Border(left=lado if cc == openpyxl_col(c1) else None,
                                right=lado if cc == openpyxl_col(c2) else None,
                                top=lado if rr == r1 else None,
                                bottom=lado if rr == r2 else None)
            if not isinstance(cel, MergedCell):
                cel.fill = fill(CARD)


def tile(ws, col, linha, rotulo, formula, fmt, sub, sub_fmt=None, neg=False):
    c1, c2 = L(col), L(col + 1)
    for rr in (linha, linha + 1, linha + 2):
        ws.merge_cells(f"{c1}{rr}:{c2}{rr}")
    put(ws, f"{c1}{linha}", rotulo, 9, True, GRAFITE, al="left")
    put(ws, f"{c1}{linha + 1}", formula, 22, True, fmt=fmt, al="left", fam=F_DISPLAY)
    put(ws, f"{c1}{linha + 2}", sub, 9, cor=GRAFITE, fmt=sub_fmt, al="left")
    ws.row_dimensions[linha].height = 22
    ws.row_dimensions[linha + 1].height = 32
    ws.row_dimensions[linha + 2].height = 20
    card(ws, c1, c2, linha, linha + 2)
    for rr in (linha, linha + 1, linha + 2):
        ws[f"{c1}{rr}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if neg:
        cor_condicional(ws, f"{c1}{linha + 1}", f"AND(ISNUMBER({c1}{linha + 1}),{c1}{linha + 1}<0)")
    return f"{c1}{linha + 1}"


def novo(tipo, titulo, fmt):
    ch = BarChart() if tipo == "bar" else LineChart()
    if tipo == "bar":
        ch.type = "col"
        ch.gapWidth = 50
    estilo_grafico(ch, titulo, fmt)
    return ch


def canvas(ws):
    """Fundo papel (#F6F5F2); células com borda viram superfícies brancas."""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row + 25, max_col=ws.max_column + 6):
        for c in row:
            if isinstance(c, MergedCell) or c.fill.fill_type:
                continue
            tem_borda = any(getattr(c.border, lado).style for lado in ("left", "right", "top", "bottom"))
            c.fill = fill(CARD if tem_borda else CANVAS)
            if c.value is None and not c.has_style:
                c.font = font()


def finalizar(wb, ordem, aba_inicial):
    """Ordena abas, aplica o canvas e deixa pronto para imprimir em A4 paisagem."""
    wb._sheets = [wb[n] for n in ordem]
    wb.active = 0
    for ws in wb:
        canvas(ws)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.zoomScale = 100
        ws.sheet_view.tabSelected = ws.title == aba_inicial
